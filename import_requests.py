import requests
import base64
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Connexion à l'API et récupération des données
url = "URL_DE_L_API"  # Remplace par l'URL de l'API
params = {}  # Paramètres de la requête si nécessaire
credentials = "CREDENTIALS"  # Identifiants de connexion
encoded_credentials = base64.b64encode(credentials.encode()).decode()  # Encodage en Base64

headers = {
    "Accept": "application/json",
    "Authorization": f"Basic {encoded_credentials}"
}

# Envoi de la requête GET à l'API
response = requests.get(url, headers=headers, params=params)
response.raise_for_status()  # Vérifie que la requête a réussi

# Extraction des données en JSON
data = response.json()
resultat = data.get("resultat", [])  # Récupération des résultats sous forme de liste

# Transformation des données en DataFrame pandas
df = pd.DataFrame(resultat)

# Filtrer uniquement les offres de l'année recherchée (exemple : celles qui commencent par "O25")
df = df[df["identifiant"].str.startswith("O25", na=False)]

# Fonction pour calculer le total des montants des commandes
def format_commandes_montant(commandes):
    """Calcule le total des montants des commandes"""
    return sum(item["valeur"] for item in commandes) if isinstance(commandes, list) else 0

# Création d'une nouvelle colonne avec le total des commandes
df["TotalCommandesMontant"] = df["commandes_montant"].apply(format_commandes_montant)

# Renommage des colonnes pour une meilleure lisibilité
df = df.rename(columns={
    "identifiant": "Titre",
    "designation": "Designation",
    "client_designation": "Client",
    "dateenvoi": "DateEnvoi",
    "montant": "Montant",
    "margenette": "MargeNette"
})

# Sélection des colonnes à conserver
df = df[["Titre", "Designation", "Client", "DateEnvoi", "Montant", "MargeNette", "TotalCommandesMontant"]]

# Exportation des données vers un fichier Excel
output_file = "GBE-Offres.xlsx"
df.to_excel(output_file, index=False, sheet_name="Offres", engine='openpyxl')

# Ajout d'un tableau structuré dans le fichier Excel pour une meilleure lisibilité
wb = load_workbook(output_file)
ws = wb["Offres"]

# Définition du tableau avec un style
table = Table(displayName="TableauOffres", ref=ws.dimensions)
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True)
table.tableStyleInfo = style
ws.add_table(table)

# Sauvegarde du fichier final
wb.save(output_file)

print(f"Fichier {output_file} créé avec succès, contenant uniquement les offres de l'année recherchée ({len(df)} lignes).")
